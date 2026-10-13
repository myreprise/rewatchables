"""
Tests for step 16 and the rewatchvault.com page parser.

The fixture page is invented. It copies the site's markup -- the classes, the
nesting, the "(N)" label -- but none of its writing, so this file can sit in a
public repository.

What matters most here:
  * a name cut out of the prose goes back where it was, and its offsets point
    at it (the whole reason for re-reading the site, report 021);
  * the "names cut out" form of a page bullet equals the workbook's cell, since
    every cross-check depends on that;
  * a slug belongs to the LONGEST episode it starts with (friday vs
    friday-night-lights);
  * storing the same content twice adds no row.

No network.
"""

import json
import zlib

import pytest

from rewatchables import rewatchvault as rv
from rewatchables.pipeline import s16_load_categories as s16

PAGE = """<!doctype html><html><body><main>
<article class="archive-answer-card" id="answer-made-up-film-apex-7">
 <div class="archive-answer-context">
  <a class="archive-answer-identity" href="/archive/movies/made-up-film-1999?episode=made-up-film#entry-made-up-film-apex-7"><strong>Made Up Film</strong> <span>·</span> <span>The Rewatchables</span></a>
  <span class="archive-answer-date"><span>·</span><time datetime="2024-03-05">March 5, 2024</time></span>
 </div>
 <header class="archive-answer-banner"><img class="archive-icon" src="/assets/category-awards/apex-0a1b2c3d.svg"/><h1>Apex Mountain</h1></header>
 <p class="archive-answer archive-notebook-answer">The  invented   answer, by <a class="archive-person-mention" href="/archive/people/film-person:abc">Pat Example</a></p>
 <details class="archive-discussion">
  <summary>In the discussion <span>(3)</span></summary>
  <div class="archive-transcript-reference"><span class="archive-transcript-time">Approx. transcript time · 01:02:03</span><a href="https://example.test/transcript">Read transcript</a></div>
  <ul>
   <li>An opening remark about <a class="archive-movie-mention" href="/archive/movies/other-film-1990">Other Film</a>.</li>
   <li>Before the climb, <a class="archive-person-mention" href="/archive/people/film-person:abc">Pat Example</a> refuses
       the rope; <a class="archive-person-mention" href="/archive/people/film-person:def">Lee Sample</a> laughs. <!-- a comment --></li>
   <li><a class="archive-person-mention" href="/archive/people/film-person:abc">Pat Example</a>, who plays the guide, gets the nod. <span class="archive-speaker">— <a href="/archive/people/host-one">Host One</a></span></li>
  </ul>
 </details>
 <details class="archive-discussion archive-podcast-choices">
  <summary>Podcast choices</summary>
  <dl>
   <div><dt>Speaker not identified</dt><dd><a class="archive-person-mention" href="/archive/people/film-person:def">Lee Sample</a></dd></div>
   <div><dt><a href="/archive/people/host-one">Host One</a></dt><dd>The <a class="archive-movie-mention" href="/archive/movies/other-film-1990">Other Film</a> cameo</dd></div>
  </dl>
 </details>
</article></main></body></html>"""


@pytest.fixture
def page():
    return rv.parse_article(rv.article_html(PAGE))


class TestParse:
    def test_entry_fields(self, page):
        assert page["discussion_slug"] == "made-up-film-apex-7"
        assert page["episode_slug"] == "made-up-film"
        assert page["film_slug"] == "made-up-film-1999"
        assert page["film_title"] == "Made Up Film"
        assert page["air_date"] == "2024-03-05"
        assert page["icon_family"] == "apex"
        assert page["category_slug"] == "apex"
        assert page["banner"] == "Apex Mountain"
        assert page["answer"] == "The invented answer, by Pat Example"
        assert page["bullet_count"] == 3
        assert page["transcript_seconds"] == 3723
        assert page["transcript_url"] == "https://example.test/transcript"

    def test_names_go_back_into_the_prose(self, page):
        b = page["bullets"][1]
        assert b["text"] == ("Before the climb, Pat Example refuses the rope; "
                             "Lee Sample laughs.")
        assert [m["name"] for m in b["mentions"]] == ["Pat Example", "Lee Sample"]

    def test_offsets_point_at_the_names(self, page):
        answer = {"text": page["answer"], "mentions": page["answer_mentions"]}
        for inline in [answer, *page["bullets"], *page["choices"]]:
            for m in inline["mentions"]:
                assert inline["text"][m["start"]:m["end"]] == m["name"]

    def test_film_links_are_mentions_too(self, page):
        b = page["bullets"][0]
        assert b["text"] == "An opening remark about Other Film."
        assert b["mentions"] == [{"kind": "movie", "slug": "other-film-1990", "name": "Other Film",
                                  "start": 24, "end": 34}]

    def test_an_answer_keeps_its_links(self, page):
        assert page["answer_mentions"][0]["slug"] == "film-person:abc"
        assert not rv.looks_excised(page["answer"])

    def test_the_hosts_picks_are_their_own_list_not_bullets(self, page):
        assert len(page["bullets"]) == 3
        assert page["choices"] == [
            {"speaker": {"person_slug": None, "name": "Speaker not identified"},
             "text": "Lee Sample",
             "mentions": [{"kind": "person", "slug": "film-person:def", "name": "Lee Sample",
                           "start": 0, "end": 10}]},
            {"speaker": {"person_slug": "host-one", "name": "Host One"},
             "text": "The Other Film cameo",
             "mentions": [{"kind": "movie", "slug": "other-film-1990", "name": "Other Film",
                           "start": 4, "end": 14}]},
        ]

    def test_the_picks_block_is_never_taken_for_the_discussion(self):
        """Both blocks carry details.archive-discussion; the picks may come first."""
        picks = PAGE[PAGE.index(' <details class="archive-discussion archive-podcast-choices">'):
                     PAGE.index("</article>")]
        swapped = PAGE.replace(picks, "").replace(' <details class="archive-discussion">',
                                                  picks + ' <details class="archive-discussion">')
        assert swapped.index("archive-podcast-choices") < swapped.index("(3)")
        assert rv.parse_article(swapped)["bullet_count"] == 3

    def test_a_bullet_can_open_with_a_name(self, page):
        b = page["bullets"][2]
        assert b["text"].startswith("Pat Example, who plays")
        assert b["mentions"][0]["start"] == 0

    def test_the_speaker_is_kept_apart_from_the_prose(self, page):
        b = page["bullets"][2]
        assert b["speakers"] == [{"person_slug": "host-one", "name": "Host One"}]
        assert "Host One" not in b["text"]
        assert page["bullets"][0]["speakers"] == []

    def test_the_bullets_match_their_count(self, page):
        assert len(page["bullets"]) == page["bullet_count"]

    def test_no_restored_bullet_reads_excised(self, page):
        assert not any(rv.looks_excised(b["text"]) for b in page["bullets"])

    def test_without_mentions_is_what_the_workbook_scraper_saw(self, page):
        # The workbook keeps the double space where a name was; clean() folds it.
        assert rv.without_mentions(page["bullets"][1]) == rv.clean(
            "Before the climb,  refuses the rope;  laughs.")
        assert rv.without_mentions(page["bullets"][2]) == ", who plays the guide, gets the nod."

    def test_a_page_with_no_discussion_is_an_entry_without_one(self):
        """Batch 1 has one: an answer and a transcript time, no discussion block."""
        start = PAGE.index('<details class="archive-discussion">')
        end = PAGE.index("</details>", start) + len("</details>")
        ref = ('<div class="archive-transcript-reference"><span class="archive-transcript-time">'
               'Approx. transcript time · 01:04:19</span><a href="https://example.test/t">x</a></div>')
        page = rv.parse_article(PAGE[:start] + ref + PAGE[end:])
        assert page["has_discussion"] is False
        assert (page["bullets"], page["bullet_count"]) == ([], None)
        assert page["transcript_seconds"] == 3859
        assert page["answer"] == "The invented answer, by Pat Example"

    def test_a_page_with_no_identity_is_an_error(self):
        html = PAGE.replace('class="archive-answer-identity"', 'class="something-else"')
        with pytest.raises(rv.ParseError):
            rv.parse_article(html)

    def test_the_source_text_is_the_yardstick(self, page):
        """Parsed bullets equal the page's own text: nothing was dropped."""
        texts = rv.source_bullet_texts(PAGE)
        assert texts == [rv.clean(b["text"]) for b in page["bullets"]]
        assert "Host One" not in texts[2]

    def test_a_gap_in_the_sites_own_prose_is_a_quirk_not_a_failure(self, tmp_path):
        """Batch 3: "seasons through the 1960s ." is how the page has it."""
        html = PAGE.replace("<li>An opening remark about", "<li>Through the 1960s . An opening remark about")
        conn = s16.connect(tmp_path / "t.db")
        TestPlanAndCheck()._index(conn)
        load = s16.start_load(conn, "web", "pilot", "x", "y", 1, ["made-up-film-apex-7"])
        art = rv.article_html(html)
        s16.store(conn, "raw_category_pages", load, "made-up-film-apex-7", s16._sha(art),
                  {"url": "u", "http_status": 200, "article_html": zlib.compress(art.encode())})
        result = s16.check_batch(conn, "pilot", ["made-up-film-apex-7"], out_dir=tmp_path)
        conn.close()
        assert result["passed"]
        assert result["problems"]["source quirk"] == 1
        assert result["problems"]["reads excised"] == 0

    def test_a_decimal_is_not_a_cut_out_name(self):
        assert not rv.looks_excised("Ted Williams hit .406 with a .553 on-base")
        assert not rv.looks_excised("this is a game ... of focus")
        assert not rv.looks_excised("having written ...And Justice for All")
        assert rv.looks_excised("makeup, , would not repeat")
        assert rv.looks_excised(", who played the hitchhiker")

    def test_a_page_with_no_article_is_an_error(self):
        with pytest.raises(rv.ParseError):
            rv.article_html("<html><body>Not found</body></html>")

    def test_no_answer_is_none(self):
        html = PAGE[:PAGE.index('<p class="archive-answer')] + PAGE[PAGE.index("</p>") + 4:]
        assert rv.parse_article(html)["answer"] is None


class TestSlugs:
    def test_the_longest_episode_wins(self):
        owners = rv.owners(["friday-night-lights-apex-12", "friday-apex-3"],
                           ["friday", "friday-night-lights"])
        assert owners == {"friday-night-lights-apex-12": "friday-night-lights",
                          "friday-apex-3": "friday"}

    def test_an_unknown_episode_has_no_owner(self):
        assert rv.owners(["moneyball-apex-1"], ["money"]) == {"moneyball-apex-1": None}

    def test_category_slug(self):
        assert rv.category_slug("the-re-heat-recasting-couch-0241", "the-re-heat") == "recasting-couch"
        assert rv.category_slug("heat-apex-1", "the-re-heat") is None

    def test_sitemap_slugs(self):
        xml = ("<urlset><url><loc>https://rewatchvault.com/archive/discussions/a-apex-1</loc></url>"
               "<url><loc>https://rewatchvault.com/archive/movies/a-1999</loc></url>"
               "<url><loc>https://rewatchvault.com/archive/discussions/a-apex-1</loc></url></urlset>")
        assert rv.sitemap_discussion_slugs(xml) == ["a-apex-1"]

    def test_transcript_seconds(self):
        assert rv.transcript_seconds("Approx. transcript time · 00:25:51") == 1551
        assert rv.transcript_seconds("no time") is None


def _pairs():
    """A workbook row for the fixture page, as the scraper laid it out."""
    return [
        ["archive-answer-identity href",
         "https://rewatchvault.com/archive/movies/made-up-film-1999?episode=made-up-film#entry-x"],
        ["archive-answer-identity", "Made Up Film"],
        ["archive-answer-date 2", "March 5, 2024"],
        ["archive-answer-poster href", "https://rewatchvault.com/archive/discussions/made-up-film-apex-7"],
        ["archive-icon src", "https://rewatchvault.com/assets/category-awards/apex-0a1b2c3d.svg"],
        ["archive-answer-banner", "Apex Mountain"],
        ["archive-answer", "The invented answer, by "],
        ["archive-discussion 2", "(3)"],
        ["archive-transcript-time", "Approx. transcript time · 01:02:03"],
        ["archive-transcript-reference href", "https://example.test/transcript"],
        ["archive-discussion 3", "An opening remark about ."],
        ["archive-discussion 4", "Before the climb,  refuses the rope;  laughs."],
        ["archive-person-mention", "Pat Example"],
        ["archive-discussion 5", "Speaker not identified"],
        ["archive-discussion 6", "Podcast choices"],
    ]


class TestCompare:
    def test_a_faithful_row_has_no_disagreements(self, page):
        diffs, tallies = s16.compare(page, _pairs())
        assert diffs == []
        # The third cell is a speaker label from the picks block, which the
        # scraper's discussion columns also caught.
        assert tallies["wb_bullets"] == 4 and tallies["wb_bullets_matched"] == 4
        # Lee Sample was in the prose but not in this row's mention columns.
        assert tallies["page_mentions_not_in_wb"] == 1

    def test_a_changed_answer_is_reported(self, page):
        pairs = [[h, "Another answer" if h == "archive-answer" else v] for h, v in _pairs()]
        diffs, _ = s16.compare(page, pairs)
        assert diffs == [("answer", "Another answer", "The invented answer, by")]

    def test_workbook_row_keys(self):
        assert s16.workbook_row(_pairs()) == ("made-up-film-apex-7", "made-up-film", "2024-03-05")


@pytest.fixture
def conn(tmp_path):
    c = s16.connect(tmp_path / "t.db")
    yield c
    c.close()


class TestStore:
    def _store(self, conn, load_id, html):
        return s16.store(conn, "raw_category_pages", load_id, "made-up-film-apex-7",
                         s16._sha(html), {"url": "u", "http_status": 200,
                                          "article_html": zlib.compress(html.encode())})

    def test_same_content_twice_is_one_row(self, conn):
        a = s16.start_load(conn, "web", "pilot", "x", "y", 1)
        b = s16.start_load(conn, "web", "pilot", "x", "y", 1)
        assert self._store(conn, a, "<article/>") == "new"
        assert self._store(conn, b, "<article/>") == "unchanged"
        rows = conn.execute("SELECT load_id, last_seen_load_id FROM raw_category_pages").fetchall()
        assert rows == [(a, b)]

    def test_changed_content_is_a_new_version(self, conn):
        a = s16.start_load(conn, "web", "pilot", "x", "y", 1)
        self._store(conn, a, "<article>1</article>")
        assert self._store(conn, a, "<article>2</article>") == "new"
        assert conn.execute("SELECT COUNT(*) FROM raw_category_pages").fetchone()[0] == 2

    def test_check_reads_the_newest_version(self, conn):
        a = s16.start_load(conn, "web", "pilot", "x", "y", 1)
        self._store(conn, a, "<article>old</article>")
        self._store(conn, a, rv.article_html(PAGE))
        assert s16.latest_pages(conn, ["made-up-film-apex-7"]) == {
            "made-up-film-apex-7": rv.article_html(PAGE)}


class TestPlanAndCheck:
    def _index(self, conn):
        load = s16.start_load(conn, "workbook", None, "f.xlsx", "h", 1)
        text = json.dumps(_pairs())
        s16.store(conn, "raw_category_workbook_rows", load, "made-up-film-apex-7", s16._sha(text),
                  {"source_file": "f.xlsx", "source_row_number": 2, "episode_slug": "made-up-film",
                   "air_date": "2024-03-05", "payload": zlib.compress(text.encode())})
        s16.finish_load(conn, load, {"new": 1, "unchanged": 0}, [], "success")
        return s16.workbook_index(conn)

    def test_pages_go_to_their_air_year_and_strays_to_the_last_batch(self, conn):
        batches, owner = s16.plan(["made-up-film-apex-7", "made-up-film-scene-8", "c54-apex-1"],
                                  self._index(conn))
        assert batches["5"] == ["made-up-film-apex-7", "made-up-film-scene-8"]   # 2024
        assert batches[s16.LAST_BATCH] == ["c54-apex-1"]
        assert owner["c54-apex-1"] is None
        assert "pilot" not in batches

    def test_a_clean_batch_passes(self, conn, tmp_path):
        index = self._index(conn)
        load = s16.start_load(conn, "web", "pilot", "x", "y", 1)
        art = rv.article_html(PAGE)
        s16.store(conn, "raw_category_pages", load, "made-up-film-apex-7", s16._sha(art),
                  {"url": "u", "http_status": 200, "article_html": zlib.compress(art.encode())})
        result = s16.check_batch(conn, "pilot", ["made-up-film-apex-7"], index, out_dir=tmp_path)
        assert result["passed"], result["summary"].read_text(encoding="utf-8")
        assert result["tallies"]["wb_bullets_matched"] == 4

    def test_a_page_not_stored_fails_the_gate(self, conn, tmp_path):
        index = self._index(conn)
        result = s16.check_batch(conn, "pilot", ["made-up-film-apex-7"], index, out_dir=tmp_path)
        assert not result["passed"]
        assert result["problems"]["not stored"] == 1

    def test_the_workbook_episode_beats_the_prefix(self, conn):
        """56 workbook pages have opaque slugs like c2 that start with no episode."""
        index = self._index(conn)
        index["c2"] = ("made-up-film", "2024-03-05", None)
        owner = s16.episode_owners(["c2", "made-up-film-scene-8", "c54"], index)
        assert owner == {"c2": "made-up-film", "made-up-film-scene-8": "made-up-film", "c54": None}

    def test_no_workbook_load_is_refused(self, conn):
        with pytest.raises(SystemExit):
            s16.workbook_index(conn)

    def test_the_pilot_has_one_episode_per_year(self):
        years = [int(note[:4]) for note in s16.PILOT_EPISODES.values()]
        assert sorted(years) == sorted(s16.YEAR_BATCHES)


def test_step_16_is_not_in_the_orchestrated_steps():
    """No --all run should start a crawl; batches are run one at a time, by hand."""
    from rewatchables.pipeline import orchestrate
    assert not any(step[1] == "s16_load_categories" for step in orchestrate.STEPS)
    assert "s16_load_categories" in orchestrate.MANUAL_STEPS


def test_the_tables_are_in_the_raw_zone(tmp_path):
    conn = s16.connect(tmp_path / "t.db")
    names = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    conn.close()
    assert {"raw_category_loads", "raw_category_workbook_rows", "raw_category_pages"} <= names


def test_export_writes_one_row_per_entry_bullet_and_pick(tmp_path):
    from openpyxl import load_workbook

    conn = s16.connect(tmp_path / "t.db")
    TestPlanAndCheck()._index(conn)
    load = s16.start_load(conn, "web", "pilot", "x", "y", 1, ["made-up-film-apex-7"])
    art = rv.article_html(PAGE)
    s16.store(conn, "raw_category_pages", load, "made-up-film-apex-7", s16._sha(art),
              {"url": "u", "http_status": 200, "article_html": zlib.compress(art.encode())})
    path, entries, bullets, picks = s16.export_batch(conn, "pilot", out_dir=tmp_path)
    conn.close()
    assert (entries, bullets, picks) == (1, 3, 2)
    wb = load_workbook(path, read_only=True)
    discussion = list(wb["Discussion"].iter_rows(values_only=True))
    assert discussion[2][4] == "Before the climb, Pat Example refuses the rope; Lee Sample laughs."
    assert discussion[2][6] == "Pat Example; Lee Sample"
    wb.close()
